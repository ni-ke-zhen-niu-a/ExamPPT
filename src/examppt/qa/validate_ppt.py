from pptx import Presentation
from pathlib import Path
import argparse, json, re

ROOT = Path(__file__).resolve().parent

def load(p): return json.loads(Path(p).read_text(encoding="utf-8"))

def validate(pptx, manifest_path, rules_path):
    m, rules = load(manifest_path), load(rules_path)
    prs = Presentation(str(pptx))
    S, L = rules["slide"], rules["layout"]
    tol = 0.015
    issues, order, details = [], [], []

    if abs(prs.slide_width/914400 - S["width_in"]) > tol or abs(prs.slide_height/914400 - S["height_in"]) > tol:
        issues.append("slide size mismatch")

    for si, sl in enumerate(prs.slides, 1):
        qs = []
        for sh in sl.shapes:
            mm = re.fullmatch(r"Q(\d{3})_CONTENT", sh.name or "")
            if not mm: continue
            q = int(mm.group(1))
            qs.append((sh.top, q, sh))
            if abs(sh.width/914400 - L["question_width_in"]) > tol:
                issues.append(f"slide {si} Q{q}: width mismatch")
            if sh.left < 0 or sh.top < 0 or sh.left+sh.width > prs.slide_width or sh.top+sh.height > prs.slide_height:
                issues.append(f"slide {si} Q{q}: out of slide bounds")
            bottom_margin = S["height_in"] - (sh.top+sh.height)/914400
            if bottom_margin < L["min_bottom_margin_in"] - tol:
                issues.append(f"slide {si} Q{q}: bottom margin {bottom_margin:.3f}in")
        qs.sort(key=lambda x:x[0])
        qnums=[x[1] for x in qs]
        order.extend(qnums)
        header=[sh.text for sh in sl.shapes if getattr(sh,"name","")=="HEADER_RANGE" and hasattr(sh,"text")]
        details.append({"slide":si,"questions":qnums,"header":header[0] if header else ""})

    expected=list(range(1,m["question_count"]+1))
    if order != expected:
        issues.append(f"question order mismatch: {order}")
    if len(set(order)) != len(order):
        issues.append("duplicate question detected")

    starts=set(m.get("section_break_before",[]))
    for b in starts:
        found=[d for d in details if b in d["questions"]]
        if not found or found[0]["questions"][0] != b:
            issues.append(f"Q{b}: section must start on new slide")

    report={"status":"PASS" if not issues else "MANUAL_REQUIRED","pptx":str(pptx),"slides":len(prs.slides),"questions":order,"details":details,"issues":issues}
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("--pptx",required=True)
    ap.add_argument("--manifest",required=True)
    ap.add_argument("--rules",default=str(ROOT/"layout_rules.json"))
    ap.add_argument("--out")
    a=ap.parse_args()
    rep=validate(a.pptx,a.manifest,a.rules)
    if a.out: Path(a.out).write_text(json.dumps(rep,ensure_ascii=False,indent=2),encoding="utf-8")
