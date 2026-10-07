from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image
from io import BytesIO
from pathlib import Path
import argparse, json, numpy as np

ROOT = Path(__file__).resolve().parent

def load_json(p):
    return json.loads(Path(p).read_text(encoding="utf-8"))

def pictures(slide):
    return [s for s in slide.shapes if s.shape_type == 13]

def compact_blank_rows(im, keep=10, minrun=18, threshold=245):
    a = np.array(im)
    ink = np.any(a < threshold, axis=2)
    row = ink.any(axis=1)
    chunks, i = [], 0
    while i < len(row):
        j = i + 1
        while j < len(row) and row[j] == row[i]:
            j += 1
        if row[i]:
            chunks.append(a[i:j])
        else:
            length = j - i
            kept = length if length < minrun else min(length, keep)
            if kept:
                chunks.append(a[i:i+kept])
        i = j
    return Image.fromarray(np.concatenate(chunks, axis=0))

def stack_images(images, gap_px=10):
    width = max(im.width for im in images)
    height = sum(im.height for im in images) + gap_px * (len(images)-1)
    out = Image.new("RGB", (width, height), "white")
    y = 0
    for im in images:
        out.paste(im, (0, y))
        y += im.height + gap_px
    return out

def section_title(manifest, q):
    for s in manifest["sections"]:
        if s["start"] <= q <= s["end"]:
            return s["title"]
    return ""

def qlabel(qs):
    return f"第 {qs[0]} 题" if len(qs) == 1 else f"第 {qs[0]}–{qs[-1]} 题"

def build(manifest_path, rules_path):
    m = load_json(manifest_path)
    rules = load_json(rules_path)
    src = Path(m["source_pptx"])
    out = Path(m["output_pptx"])
    qa_path = Path(m["qa_json"])
    assets = ROOT / "assets" / Path(manifest_path).stem
    assets.mkdir(parents=True, exist_ok=True)

    src_prs = Presentation(str(src))
    qimgs = {}
    violations = []

    for q in range(1, m["question_count"] + 1):
        refs = m["question_map"].get(str(q))
        if not refs:
            violations.append(f"Q{q}: missing mapping")
            continue
        parts = []
        for part_idx, (slide_no, pic_no) in enumerate(refs, 1):
            pics = pictures(src_prs.slides[slide_no - 1])
            if pic_no > len(pics):
                violations.append(f"Q{q}: source picture {slide_no}/{pic_no} missing")
                continue
            im = Image.open(BytesIO(pics[pic_no - 1].image.blob)).convert("RGB")
            proc = m.get("processing", {}).get(str(q), {})
            crop = proc.get("crop_bottom_px", {}).get(str(part_idx))
            if crop:
                im = im.crop((0, 0, im.width, min(crop, im.height)))
            if part_idx in proc.get("compact_parts", []):
                im = compact_blank_rows(im)
            parts.append(im)
        if not parts:
            continue
        comp = stack_images(parts) if len(parts) > 1 else parts[0]
        p = assets / f"q{q:02d}.png"
        comp.save(p)
        qimgs[q] = {"path": p, "ratio": comp.height / comp.width}

    L = rules["layout"]
    S = rules["slide"]
    content_limit = S["height_in"] - L["bottom_in"]
    slides, current, used = [], [], 0.0
    breaks = set(m.get("section_break_before", []))

    for q in range(1, m["question_count"] + 1):
        if q not in qimgs:
            continue
        if q in breaks and current:
            slides.append(current)
            current, used = [], 0.0
        h = L["question_width_in"] * qimgs[q]["ratio"]
        max_h = content_limit - L["content_top_in"]
        if h > max_h:
            violations.append(f"Q{q}: intrinsic height {h:.3f}in exceeds content area {max_h:.3f}in; MANUAL_REQUIRED")
        needed = h if not current else L["gap_in"] + h
        if current and L["content_top_in"] + used + needed > content_limit:
            slides.append(current)
            current, used = [], 0.0
        current.append(q)
        used += h if len(current) == 1 else L["gap_in"] + h
    if current:
        slides.append(current)

    prs = Presentation()
    prs.slide_width = Inches(S["width_in"])
    prs.slide_height = Inches(S["height_in"])
    blank = prs.slide_layouts[6]
    qa = []

    for idx, qs in enumerate(slides, 1):
        sl = prs.slides.add_slide(blank)
        title = sl.shapes.add_textbox(Inches(0.35), Inches(0.10), Inches(7.9), Inches(0.40))
        title.name = "HEADER_TITLE"
        p = title.text_frame.paragraphs[0]
        r = p.add_run(); r.text = m["name"]
        r.font.name = rules["header"]["title_font"]; r.font.size = Pt(rules["header"]["title_pt"]); r.font.bold = True
        r.font.color.rgb = RGBColor(35,35,35)

        right = sl.shapes.add_textbox(Inches(8.1), Inches(0.11), Inches(4.9), Inches(0.36))
        right.name = "HEADER_RANGE"
        p = right.text_frame.paragraphs[0]
        r = p.add_run(); r.text = section_title(m, qs[0])
        r.font.name = rules["header"]["title_font"]; r.font.size = Pt(rules["header"]["range_pt"]); r.font.bold = True
        r.font.color.rgb = RGBColor(70,70,70); p.alignment = PP_ALIGN.RIGHT

        line = sl.shapes.add_shape(1, Inches(0.35), Inches(rules["header"]["divider_y_in"]), Inches(12.65), Inches(0.01))
        line.name = "HEADER_DIVIDER"; line.fill.solid(); line.fill.fore_color.rgb = RGBColor(225,225,225); line.line.color.rgb = RGBColor(225,225,225)

        y = L["content_top_in"]
        entries = []
        for q in qs:
            h = L["question_width_in"] * qimgs[q]["ratio"]
            shp = sl.shapes.add_picture(str(qimgs[q]["path"]), Inches(L["left_in"]), Inches(y), width=Inches(L["question_width_in"]))
            shp.name = f"Q{q:03d}_CONTENT"
            entries.append({"q": q, "top": round(y,3), "bottom": round(y+h,3), "height": round(h,3)})
            y += h + L["gap_in"]
        bottom_margin = S["height_in"] - entries[-1]["bottom"]
        if bottom_margin < L["min_bottom_margin_in"]:
            violations.append(f"Slide {idx}: bottom margin {bottom_margin:.3f}in below minimum")
        qa.append({"slide": idx, "questions": qs, "bottom_margin": round(bottom_margin,3), "entries": entries})

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))

    seen = [q for s in qa for q in s["questions"]]
    expected = list(range(1, m["question_count"] + 1))
    if seen != expected:
        violations.append("Question order/coverage mismatch")
    for b in breaks:
        for s in qa:
            if b in s["questions"] and s["questions"][0] != b:
                violations.append(f"Q{b}: section does not start on a new slide")

    report = {
        "status": "PASS" if not violations else "MANUAL_REQUIRED",
        "output": str(out),
        "slide_count": len(qa),
        "question_count": m["question_count"],
        "fixed_question_width_in": L["question_width_in"],
        "questions": seen,
        "slides": qa,
        "violations": violations
    }
    qa_path.parent.mkdir(parents=True, exist_ok=True)
    qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--rules", default=str(ROOT / "layout_rules.json"))
    args = ap.parse_args()
    build(args.manifest, args.rules)
